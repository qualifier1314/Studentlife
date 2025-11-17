import os
import csv
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN


def add_left_image_right_text_slide(prs, title, image_path, right_bullets):
    """添加左侧图片、右侧文字要点的双栏页面。"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9.5), Inches(0.8))
    tf = title_box.text_frame
    tf.text = title
    tf.paragraphs[0].font.size = Pt(28)
    tf.paragraphs[0].font.bold = True

    # Left image
    if os.path.exists(image_path):
        slide.shapes.add_picture(image_path, Inches(0.5), Inches(1.2), width=Inches(6.0))

    # Right text bullets（加宽以容纳更多解释）
    box = slide.shapes.add_textbox(Inches(7.0), Inches(1.2), Inches(3.8), Inches(6.5))
    t = box.text_frame
    t.word_wrap = True
    for i, b in enumerate(right_bullets):
        p = t.add_paragraph() if i > 0 else t.paragraphs[0]
        p.text = b
        p.level = 0
        p.font.size = Pt(16)


def add_bullets_slide(prs, title, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
    tf = title_box.text_frame
    tf.text = title
    tf.paragraphs[0].font.size = Pt(28)
    tf.paragraphs[0].font.bold = True

    # Bullets
    box = slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(8.5), Inches(6.5))
    t = box.text_frame
    t.word_wrap = True
    for i, b in enumerate(bullets):
        p = t.add_paragraph() if i > 0 else t.paragraphs[0]
        p.text = b
        p.level = 0
        p.font.size = Pt(18)


def ensure_comm_features_figure(out_dir, fig_dir):
    """返回通信特征分布图路径；若存在直接使用，不再内联生成。"""
    os.makedirs(fig_dir, exist_ok=True)
    fig_path = os.path.join(fig_dir, 'Communication_features.png')
    return fig_path if os.path.exists(fig_path) else None


def smart_save(prs, out_dir, preferred_name='report.pptx'):
    """避免被占用导致无法保存：尝试报告文件名递增保存。"""
    base = os.path.join(out_dir, preferred_name)
    try:
        prs.save(base)
        print(f"[Saved] {base}")
        return base
    except PermissionError:
        # 递增寻找可写文件名
        for i in range(2, 10):
            cand = os.path.join(out_dir, f"report_v{i}.pptx")
            try:
                prs.save(cand)
                print(f"[Saved] {cand} (原文件占用或锁定，已改名保存)")
                return cand
            except PermissionError:
                continue
        raise


def main():
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    out_dir = os.path.join(base_dir, 'outputs', 'static')
    fig_dir = os.path.join(out_dir, 'figures')

    prs = Presentation()

    # Title slide
    title_slide = prs.slides.add_slide(prs.slide_layouts[6])
    title_box = title_slide.shapes.add_textbox(Inches(0.5), Inches(2.0), Inches(9), Inches(1.5))
    tf = title_box.text_frame
    tf.text = '学生行为多层网络构建与可视化报告'
    tf.paragraphs[0].font.size = Pt(36)
    tf.paragraphs[0].font.bold = True
    sub = title_slide.shapes.add_textbox(Inches(0.5), Inches(3.3), Inches(9), Inches(1))
    st = sub.text_frame
    st.text = '静态综合单层 + 多层可行性 + 关键节点识别'
    st.paragraphs[0].font.size = Pt(18)

    # 建模设定（静态、存在即为边；communication 为节点属性）
    add_bullets_slide(
        prs,
        '建模设定（静态与通信属性）',
        [
            '不做时间序列：全期静态网络',
            '存在即为边：任意时刻有交互即连边（导出无权版）',
            '同时导出加权版：强度为次数/时长等累计',
            'communication 不做学生-学生层，作为节点属性（features.csv）'
        ]
    )

    # Physical（左图右文）
    add_left_image_right_text_slide(
        prs,
        'Physical（物理接近+共同就餐）静态综合层',
        os.path.join(fig_dir, 'Physical.png'),
        [
            '为什么：反映线下真实接触与共同活动的社会联系',
            'sensing 数据集：接近事件（uid_i, uid_j, ts, 距离/强度）',
            'dining 数据集：同地同时间就餐（uid, ts, 地点/餐厅）',
            '构边：累计接触时长 + 共餐次数；各源先归一化再合并',
            '静态规则：全期存在即为边；导出无权与加权版本'
        ]
    )

    # Behavior（左图右文）
    add_left_image_right_text_slide(
        prs,
        'Behavior（日历共现+应用相似度）静态综合层',
        os.path.join(fig_dir, 'Behavior.png'),
        [
            '为什么：捕捉行为相似与协作关系（同一活动、相近习惯）',
            'calendar：事件起止与参与者，统计两两共现次数',
            'app_usage：应用使用频谱（包名×时间），计算余弦相似度',
            '构边：共现计数 + 相似度加权；归一化后合并',
            '说明：当前静态层使用 calendar+app_usage（rawaccfeat 可选）'
        ]
    )

    # Education（左图右文）
    add_left_image_right_text_slide(
        prs,
        'Education（同课共修）静态层',
        os.path.join(fig_dir, 'Education.png'),
        [
            '为什么：课程共修形成稳定学习社群与协作网络',
            '数据集：education/class.csv（uid + 课程代码列表）',
            '构边：共享课程计数作为边权；共享课程数>0 即连边',
            '说明：可构学生-课程二分图，本报告采用共修投影',
            '静态规则：全期存在即为边；导出无权/加权'
        ]
    )

    # Communication（节点属性）左图右文：使用已生成的特征分布图
    comm_fig = ensure_comm_features_figure(out_dir, fig_dir)
    if comm_fig:
        add_left_image_right_text_slide(
            prs,
            'Communication（节点属性层）',
            comm_fig,
            [
                '为什么：用通信活跃度刻画节点属性，辅助解释结构',
                '数据来源：call_log/call_log_*.csv；sms/sms_*.csv（匿名号码/地址）',
                '特征：calls_total、sms_total、comm_freq_day/week、unique_contacts、entropy',
                '限制：无法识别对方学生，不构边；仅保留节点属性',
                '左图：unique_contacts 直方图 + 每周频次散点'
            ]
        )

    # 多层网络（三层）左图右文
    add_left_image_right_text_slide(
        prs,
        '多层网络（三层：Physical + Behavior + Education）',
        os.path.join(fig_dir, 'Multiplex3D.png'),
        [
            '为什么：融合多维关系提升鲁棒性与解释力',
            '数据来源：各静态层（Physical、Behavior、Education）归一化合并',
            '层间耦合：omega·I（同一学生跨层相连）',
            '导出：supra_adjacency.csv；multi_layer_edges.csv',
            '可视化：Multiplex3D.png（三层示意图)'
        ]
    )

    # Feasibility & Key Nodes
    add_bullets_slide(
        prs,
        '多层网络可行性与关键节点识别',
        [
            '多层构建链路完整：静态层、编号统一、supra矩阵与边导出已验证',
            '可视化：3D多层示意（Multiplex3D.png），可用于演示',
            '关键节点识别可行：层感知中心性、跨层参与系数、多层k-core等',
            '建议做参数敏感性分析（层间耦合omega、各层归一化方式）',
            '数据质量与节点统一影响跨层一致性，需做缺失与稳健处理'
        ]
    )

    smart_save(prs, out_dir, preferred_name='report.pptx')


if __name__ == '__main__':
    main()